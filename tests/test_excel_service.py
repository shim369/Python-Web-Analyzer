from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from web_analyzer.core.excel_service import ExcelService
from web_analyzer.models import SiteAssessment


def _make_assessment(**overrides: Any) -> SiteAssessment:
    base: dict[str, Any] = dict(
        id="1",
        job_id="job_test_123",
        date_str="8/9",
        domain_name="example.com",
        operator_name="山田 太郎",
        evaluation_result="◯",
        has_ssl="あり",
        is_always_ssl="◯",
        total_pages=5,
        max_depth=2,
        contact_fields="名前,メールアドレス",
        site_structure="Top -> Contact",
        cms_name="WordPress",
        description="コーポレートサイト",
        rejection_reason="",
        remarks="",
    )
    base.update(overrides)
    return SiteAssessment(**base)


def test_excel_export_writes_expected_cell_values(tmp_path: Path) -> None:
    """実際のビジネスロジックが返す値（◯/×/要確認）が、正しい列に書き込まれることを確認する。"""
    test_file = tmp_path / "export_test.xlsx"

    assessments = [
        _make_assessment(id="1", domain_name="ok-example.com", evaluation_result="◯"),
        _make_assessment(
            id="2",
            domain_name="ng-example.com",
            evaluation_result="×",
            rejection_reason="ページ数が多いため\nログイン・マイページ機能があるため",
        ),
        _make_assessment(id="3", domain_name="pending-example.com", evaluation_result="要確認", total_pages=None, max_depth=None),
    ]

    ExcelService.export_excel(assessments, test_file)
    assert test_file.exists()

    wb = load_workbook(str(test_file))
    ws = wb.active

    # wsがNoneでないことをmypyに明示
    assert ws is not None

    # ヘッダー行
    assert ws.cell(row=1, column=2).value == "サイト名"
    assert ws.cell(row=1, column=3).value == "調査結果"

    # データ行（A〜O列の主要どころを確認）
    assert ws.cell(row=2, column=2).value == "ok-example.com"
    assert ws.cell(row=2, column=3).value == "◯"

    assert ws.cell(row=3, column=2).value == "ng-example.com"
    assert ws.cell(row=3, column=3).value == "×"
    assert "ログイン" in str(ws.cell(row=3, column=13).value)  # M列: 不可の理由 (strキャストでoperator型エラー回避)

    assert ws.cell(row=4, column=2).value == "pending-example.com"
    assert ws.cell(row=4, column=3).value == "要確認"
    # total_pages/max_depthがNoneの場合は空文字になる (strキャストでoperator型エラー回避)
    assert str(ws.cell(row=4, column=6).value or "") == ""
    assert str(ws.cell(row=4, column=9).value or "") == ""


def test_excel_export_strips_illegal_control_characters(tmp_path: Path) -> None:
    """Excelが受け付けない制御文字が混入していても、例外にならず除去されることを確認する。"""
    test_file = tmp_path / "control_chars.xlsx"

    assessment = _make_assessment(description="備考\x0b混入テスト\x1f文字列")
    ExcelService.export_excel([assessment], test_file)

    wb = load_workbook(str(test_file))
    ws = wb.active
    assert ws is not None

    value = str(ws.cell(row=2, column=11).value or "")  # K列: 用途
    assert "\x0b" not in value
    assert "\x1f" not in value
    assert "混入テスト" in value


def test_excel_import_reads_domain_column_and_skips_header_like_rows(tmp_path: Path) -> None:
    """B列からドメイン名を読み取り、ヘッダー的な行や空行はスキップすることを確認する。"""
    test_file = tmp_path / "import_test.xlsx"

    assessments = [
        _make_assessment(id="1", domain_name="example.com"),
        _make_assessment(id="2", domain_name="another-example.co.jp"),
    ]
    ExcelService.export_excel(assessments, test_file)

    job, imported = ExcelService.import_excel(test_file, operator_name="山田 太郎", page_threshold=15)

    assert job.operator_name == "山田 太郎"
    assert job.page_threshold == 15
    assert job.status == "processing"

    assert [a.domain_name for a in imported] == ["example.com", "another-example.co.jp"]
    # インポート直後は解析前なので、各種結果フィールドは初期値のまま
    assert imported[0].evaluation_result == ""
    assert imported[0].job_id == job.id
